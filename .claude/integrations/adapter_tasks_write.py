# -*- coding: utf-8 -*-
"""WRITE ADAPTER INT-TASKS — the first governed write surface: Tasks.xlsx (sheet ԱՌԱՋԱԴՐԱՆՔՆԵՐ, rows 13+, columns № · task · status · comment · due · owner).
Reached ONLY through the Action Runtime (actions.execute) after Gev's approval. Provider contract:
    precondition(op, params)  → {"exists", "object"}  current row snapshot (stale-state protection)
    find_existing(op, params) → matching open task (idempotency / reconciliation) or None
    execute(op, params)       → {"ok", "id", "row"}   the single mutation, under an exclusive file lock
    verify(op, params, res)   → {"verified", "reason", "evidence"}  independent read-back with the certified loader
Writes preserve the workbook (openpyxl keeps styles/merged cells); an Excel lock (file open) → ProviderError, nothing written."""
import sys, pathlib, datetime, re
HERE = pathlib.Path(__file__).resolve().parent
sys.path.insert(0, str(HERE)); sys.path.insert(0, str(HERE.parent / "skills"))
import health

OPS = ("tasks.create", "tasks.update", "tasks.assign", "tasks.close", "tasks.reopen", "tasks.note")
DEFAULT_STATUS = "Չսկսված"; CLOSED = "Արված"
TEAM_WORDS = {"sales team", "operations", "management", "team", "թիմ", "բաժին", "ops"}

def _x(): import executors; return executors
def _path(params):
    """Register path: explicit material parameter `register_path` (tests/other registers) → env COMMAND_CENTER_TASKS_XLSX → the canonical Tasks.xlsx.
    Never an underscore-prefixed key (the Action Runtime strips those from the material parameters)."""
    import os
    return pathlib.Path(params.get("register_path") or os.environ.get("COMMAND_CENTER_TASKS_XLSX") or _x().XLSX)
def _norm(t): return re.sub(r"\s+", " ", str(t or "").strip().lower())

def _date(v):
    if v in (None, ""): return None
    if isinstance(v, datetime.datetime): return v
    if isinstance(v, datetime.date): return datetime.datetime(v.year, v.month, v.day)
    return datetime.datetime.fromisoformat(str(v)[:19]) if "T" in str(v) else datetime.datetime.fromisoformat(str(v)[:10])

def _tasks(params): return _x().load_tasks(_path(params))
def _find(params, tid=None):
    for t in _tasks(params):
        if tid is not None and t["id"] == int(tid): return t
    return None
def _snapshot(t): return {"id": t["id"], "title": t["task"], "status": t["status"], "owner": t["owner"], "due": t["due"].isoformat() if t["due"] else None, "comment": t["comment"]} if t else None

def precondition(op, params):
    tid = params.get("target_object_id")
    if tid is None: return {"exists": False, "object": None}
    t = _find(params, tid); return {"exists": t is not None, "object": _snapshot(t)}

def find_existing(op, params):
    """Idempotency / reconciliation: the intended postcondition already present?"""
    if op == "tasks.create":
        for t in _tasks(params):
            if t["open"] and _norm(t["task"]) == _norm(params.get("title")) and _norm(t["owner"]) == _norm(params.get("owner")): return {"id": t["id"], **_snapshot(t)}
        return None
    t = _find(params, params.get("target_object_id"))
    if not t: return None
    s = _snapshot(t)
    want = _expected(op, params)
    return {"id": t["id"], **s} if all(s.get(k) == v for k, v in want.items()) else None

def _expected(op, params):
    if op == "tasks.create": return {"title": params.get("title"), "owner": params.get("owner"), "status": params.get("status") or DEFAULT_STATUS, "due": params.get("due")}
    if op == "tasks.update": return {k: params[k] for k in ("title", "status", "owner", "due", "comment") if k in params}
    if op == "tasks.assign": return {"owner": params.get("owner")}
    if op == "tasks.close": return {"status": CLOSED}
    if op == "tasks.reopen": return {"status": params.get("status") or "Ընթացքում"}
    if op == "tasks.note": return {"comment": params.get("comment")}
    return {}

def _validate(op, params):
    from actions import ProviderError
    if op not in OPS: raise ProviderError(f"unknown tasks operation {op}")
    if op in ("tasks.create", "tasks.assign"):
        owner = str(params.get("owner") or "").strip()
        if not owner or _norm(owner) in TEAM_WORDS: raise ProviderError("owner must be ONE named person (no team ownership)")
    if op == "tasks.create" and not str(params.get("title") or "").strip(): raise ProviderError("title missing")
    if op == "tasks.close" and not str(params.get("evidence") or "").strip(): raise ProviderError("closing a task requires completion evidence")
    if "due" in params and params["due"]: _date(params["due"])

def execute(op, params):
    """The mutation. Exclusive lock around load → change → save; a PermissionError (Excel has the file open) means nothing was written."""
    from actions import ProviderError
    _validate(op, params); x = _x(); path = _path(params)
    import openpyxl
    with health.locked(health.state_dir() / "tasks_write.lock"):          # lock file lives in the runtime state dir, never beside the register
        try: wb = openpyxl.load_workbook(path)
        except PermissionError as e: raise ProviderError(f"workbook locked by another program: {e}")
        ws = wb[x.SHEET]
        if op == "tasks.create":
            ids = [ws.cell(row=r, column=x.COL["id"]).value for r in range(x.FIRST_ROW, ws.max_row + 1)]; ids = [i for i in ids if isinstance(i, int)]
            new_id = (max(ids) + 1) if ids else 1
            from openpyxl.cell.cell import MergedCell
            def free(r): return all(not isinstance(ws.cell(row=r, column=c), MergedCell) for c in x.COL.values()) and not isinstance(ws.cell(row=r, column=x.COL["id"]).value, int) and not ws.cell(row=r, column=x.COL["task"]).value
            row = next((r for r in range(x.FIRST_ROW, ws.max_row + 2) if free(r)), None)
            if row is None:
                row = ws.max_row + 1
                while not free(row): row += 1
            ws.cell(row=row, column=x.COL["id"], value=new_id); ws.cell(row=row, column=x.COL["task"], value=params["title"]); ws.cell(row=row, column=x.COL["status"], value=params.get("status") or DEFAULT_STATUS)
            ws.cell(row=row, column=x.COL["comment"], value=params.get("comment") or ""); ws.cell(row=row, column=x.COL["owner"], value=params["owner"])
            if params.get("due"): c = ws.cell(row=row, column=x.COL["due"], value=_date(params["due"])); c.number_format = "yyyy-mm-dd"
            tid = new_id
        else:
            tid = int(params["target_object_id"]); row = next((r for r in range(x.FIRST_ROW, ws.max_row + 1) if ws.cell(row=r, column=x.COL["id"]).value == tid), None)
            if row is None: raise ProviderError(f"task {tid} not found")
            if op == "tasks.assign": ws.cell(row=row, column=x.COL["owner"], value=params["owner"])
            elif op == "tasks.close": ws.cell(row=row, column=x.COL["status"], value=CLOSED); ws.cell(row=row, column=x.COL["comment"], value=(str(ws.cell(row=row, column=x.COL["comment"]).value or "") + f" · closed: {params['evidence']}").strip(" ·"))
            elif op == "tasks.reopen": ws.cell(row=row, column=x.COL["status"], value=params.get("status") or "Ընթացքում")
            elif op == "tasks.note": ws.cell(row=row, column=x.COL["comment"], value=params["comment"])
            elif op == "tasks.update":
                for k, col in (("title", "task"), ("status", "status"), ("owner", "owner"), ("comment", "comment")):
                    if k in params: ws.cell(row=row, column=x.COL[col], value=params[k])
                if "due" in params: c = ws.cell(row=row, column=x.COL["due"], value=_date(params["due"]) if params["due"] else None); c.number_format = "yyyy-mm-dd"
        try: wb.save(path)
        except PermissionError as e: raise ProviderError(f"workbook locked by another program at save time — nothing persisted: {e}")
    return {"ok": True, "id": tid, "row": row}

def verify(op, params, result):
    """Independent read-back with the certified loader (never the write handle)."""
    tid = (result or {}).get("id") or params.get("target_object_id")
    t = _find(params, tid) if tid is not None else None
    if not t: return {"verified": False, "reason": f"task {tid} not found on read-back", "evidence": None}
    s = _snapshot(t); want = _expected(op, params); diffs = {k: (v, s.get(k)) for k, v in want.items() if s.get(k) != v}
    return {"verified": not diffs, "reason": ("field mismatch: " + ", ".join(f"{k} expected {a!r} got {b!r}" for k, (a, b) in diffs.items())) if diffs else "read-back matches the approved action", "evidence": {"id": t["id"], "row": t["row"], **s}}
