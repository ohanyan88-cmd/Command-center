# -*- coding: utf-8 -*-
"""LIVE DATA suite — Tasks.xlsx is a LIVE OPERATIONAL SOURCE (S09 · LIVE_REGISTER · STRUCTURE-scoped): ordinary row mutations never
invalidate the certified business-model core, never require a rebuild/certification/release, and travel to GitHub through the
lightweight LIVE DATA SYNC (skill.py sync). Structure/schema, model-source, code and policy changes still enter the release path.
Provenance/freshness of live task facts stay in the integration layer (INT-TASKS). Drift classification is deterministic and fails
closed on unknown paths. Every destructive scenario runs on temp copies / temp git repositories — the canonical Tasks.xlsx is
read-only here and its hash is asserted unchanged at module teardown."""
import unittest, json, os, sys, pathlib, tempfile, shutil, subprocess, hashlib
HERE = pathlib.Path(__file__).resolve().parent
ROOT = HERE.parent.parent
sys.path.insert(0, str(HERE))
for d in ("integrations", "skills", "business", "runtime", "policy"): sys.path.insert(0, str(ROOT / ".claude" / d))
from testing import covers
import engine, store, executors, business, tree_manifest as tm, data_sync, build_business_model as bb, bm_sources, certify_business as cb

TMP = pathlib.Path(tempfile.mkdtemp(prefix="cclive_")); engine.STATE_DIR = TMP / "state"; (TMP / "state").mkdir(parents=True, exist_ok=True); store.reset()
REAL = ROOT / "Tasks.xlsx"; REAL_SHA = hashlib.sha256(REAL.read_bytes()).hexdigest() if REAL.exists() else None
os.environ["COMMAND_CENTER_TASKS_XLSX"] = str(TMP / "guard-never-written.xlsx")          # HARD GUARD: the write adapter can never default to the real register here
def tearDownModule():
    for k in ("COMMAND_CENTER_BUSINESS_DIR", "COMMAND_CENTER_BUSINESS_ROOT"): os.environ.pop(k, None)
    business.load(force=True)
    if REAL_SHA and hashlib.sha256(REAL.read_bytes()).hexdigest() != REAL_SHA: raise AssertionError("TEST SUITE MUTATED THE REAL Tasks.xlsx — forbidden")
GOV = ("authority_checking", "approval_management", "completion_verification", "audit_logging")
BQ = "business_model_query"; DS = "data_sensitivity_awareness"
S09 = next(s for s in bm_sources.SOURCES if s["source_id"] == "S09"); SHEET = S09["structure"]["sheet"]; HDR = S09["structure"]["header_rows"]

def _copy_register(name):
    p = TMP / name; shutil.copy(REAL, p); return p
def _struct(p): return bb.xlsx_structure_sha256(p, SHEET, HDR)
def _write(op, params, path):
    import adapter_tasks_write as W
    return W.execute(op, dict(params, register_path=str(path)))
def _mutate_rows(p):
    """Every ordinary Action-Runtime mutation on a temp copy: create → assign → note → close → reopen (+ due change), through the real write adapter."""
    r = _write("tasks.create", {"title": "LIVE DATA TEST — row only", "owner": "Գև", "due": "2026-12-31", "comment": "temp copy"}, p); tid = r["id"]
    _write("tasks.assign", {"target_object_id": tid, "owner": "Արման"}, p); _write("tasks.note", {"target_object_id": tid, "comment": "note"}, p)
    _write("tasks.update", {"target_object_id": tid, "due": "2027-01-15"}, p); _write("tasks.close", {"target_object_id": tid, "evidence": "test"}, p); _write("tasks.reopen", {"target_object_id": tid}, p)
    return tid
def _rename_header(p):
    import openpyxl; wb = openpyxl.load_workbook(p); ws = wb[SHEET]
    c = ws.cell(row=executors.HDR_ROW, column=executors.COL["owner"]); c.value = (str(c.value or "") + " (renamed)"); wb.save(p)

class L01_StructureFingerprint(unittest.TestCase):
    @covers(BQ, "task_management", *GOV, kinds=("unit",))
    def test_row_mutations_keep_the_structure_fingerprint_and_the_loader_happy(self):
        p = _copy_register("rows.xlsx"); base = _struct(p); n0 = len(executors.load_tasks(p))
        tid = _mutate_rows(p)
        self.assertEqual(_struct(p), base); rows = executors.load_tasks(p); self.assertEqual(len(rows), n0 + 1)
        t = next(r for r in rows if r["id"] == tid); self.assertEqual(t["owner"], "Արման"); self.assertEqual(t["due"].isoformat()[:10], "2027-01-15")
        self.assertEqual(bb.source_fingerprint(S09, p), {"sha256": base, "scope": "STRUCTURE"})
    @covers(BQ, "source_verification", *GOV, kinds=("unit", "failure_injection"))
    def test_schema_changes_change_the_structure_fingerprint(self):
        p = _copy_register("schema.xlsx"); base = _struct(p)
        _rename_header(p); self.assertNotEqual(_struct(p), base)
        q = _copy_register("sheet.xlsx"); import openpyxl; wb = openpyxl.load_workbook(q); wb.create_sheet("Extra"); wb.save(q); self.assertNotEqual(_struct(q), base)
        r = _copy_register("nosheet.xlsx"); wb = openpyxl.load_workbook(r); wb[SHEET].title = "Renamed"; wb.save(r); self.assertNotEqual(_struct(r), base)
    @covers(BQ, *GOV, kinds=("unit",))
    def test_content_sources_still_bind_by_content(self):
        p = TMP / "doc.md"; p.write_text("a", encoding="utf-8"); s = {"source_id": "SX", "path": "doc.md"}
        a = bb.source_fingerprint(s, p); p.write_text("b", encoding="utf-8"); b = bb.source_fingerprint(s, p)
        self.assertEqual(a["scope"], "CONTENT"); self.assertNotEqual(a["sha256"], b["sha256"])
        with self.assertRaises(bb.BuildError): bb.source_fingerprint({"source_id": "SY", "fingerprint_scope": "STRUCTURE"}, p)     # STRUCTURE scope only for xlsx registers

class L02_ModelBinding(unittest.TestCase):
    """A full temp workspace (every declared source copied) built, certified and loaded through the real runtime."""
    @classmethod
    def setUpClass(cls):
        cls.root = TMP / "ws"; cls.biz = TMP / "ws_model"; cls.biz.mkdir(parents=True, exist_ok=True)
        for s in bm_sources.SOURCES:
            src = ROOT / s["path"]
            if src.exists(): (cls.root / s["path"]).parent.mkdir(parents=True, exist_ok=True); shutil.copy(src, cls.root / s["path"])
        cls.core, cls.ov, cls.snap, _ = bb.build(root=cls.root); bb.write(cls.core, cls.ov, out_dir=cls.biz)
        rec = cb.certify(cls.core, cls.ov, root=cls.root, d=cls.biz, check_git=False); (cls.biz / "certification.json").write_text(json.dumps(rec), encoding="utf-8")
        assert rec["result"] == "PASS", {k: v for k, v in rec["checks"].items() if not v["pass"]}
        os.environ["COMMAND_CENTER_BUSINESS_DIR"] = str(cls.biz); os.environ["COMMAND_CENTER_BUSINESS_ROOT"] = str(cls.root); business.load(force=True)
    @classmethod
    def tearDownClass(cls):
        for k in ("COMMAND_CENTER_BUSINESS_DIR", "COMMAND_CENTER_BUSINESS_ROOT"): os.environ.pop(k, None)
        business.load(force=True)
    def _certify(self):
        core, ov = cb.load_built(self.biz); return cb.certify(core, ov, root=self.root, d=self.biz, check_git=False)
    def _state(self): return business.model_state(business.load(force=True))
    @covers(BQ, "task_management", "daily_briefing", *GOV, kinds=("unit", "completion"))
    def test_task_row_mutations_never_invalidate_the_certified_core(self):
        p = self.root / "Tasks.xlsx"; fp0 = self.snap["S09"]["sha256"]; cf0 = self.core["sources"]["meta"]["core_fingerprint"]; sid0 = self.core["sources"]["meta"]["source_snapshot_id"]
        self.assertEqual(self.snap["S09"]["scope"], "STRUCTURE"); self.assertEqual(self.snap["S09"]["live_integration"], "INT-TASKS")
        _mutate_rows(p)
        self.assertEqual(bb.snapshot(self.root)["S09"]["sha256"], fp0)                                  # binding unchanged
        core2, ov2, snap2, _ = bb.build(root=self.root)                                                  # even a rebuild yields the same business understanding
        self.assertEqual(core2["sources"]["meta"]["core_fingerprint"], cf0); self.assertEqual(core2["sources"]["meta"]["source_snapshot_id"], sid0)
        rec = self._certify(); self.assertEqual(rec["result"], "PASS", {k: v for k, v in rec["checks"].items() if not v["pass"]})
        self.assertTrue(rec["checks"]["source_fingerprints_current"]["pass"]); self.assertTrue(rec["checks"]["extraction_invariants"]["pass"]); self.assertTrue(rec["checks"]["live_registers_bound"]["pass"])
        st = self._state(); self.assertEqual(st["state"], "CURRENT", st); self.assertEqual(st["changed"], [])
        ctx = business.context_for("deadline_management", "what is overdue", {}, None); self.assertEqual(ctx["model"]["state"], "CURRENT"); self.assertNotIn("STALE_MODEL", ctx["gaps"]); self.assertNotIn("SOURCE_CHANGED", ctx["gaps"])
    @covers(BQ, "source_verification", *GOV, kinds=("unit", "failure", "failure_injection"))
    def test_register_schema_change_still_hits_the_certification_boundary(self):
        p = self.root / "Tasks.xlsx"; keep = p.read_bytes()
        try:
            _rename_header(p)
            rec = self._certify(); self.assertEqual(rec["result"], "FAIL"); self.assertIn("S09 SOURCE_CHANGED", rec["checks"]["source_fingerprints_current"]["problems"])
            st = self._state(); self.assertEqual(st["state"], "STALE_MODEL"); self.assertEqual(st["changed"], ["S09"])
        finally: p.write_bytes(keep); business.load(force=True)
        self.assertEqual(self._state()["state"], "CURRENT")
    @covers(BQ, "source_verification", *GOV, kinds=("unit", "failure_injection"))
    def test_extracted_model_source_change_still_invalidates(self):
        p = self.root / "Actions.md"; keep = p.read_bytes()
        try:
            p.write_text(keep.decode("utf-8") + "\n- new line\n", encoding="utf-8")
            rec = self._certify(); self.assertEqual(rec["result"], "FAIL"); self.assertIn("S14 SOURCE_CHANGED", rec["checks"]["source_fingerprints_current"]["problems"])
            self.assertEqual(self._state()["state"], "STALE_MODEL")
        finally: p.write_bytes(keep); business.load(force=True)
    @covers(BQ, *GOV, kinds=("unit", "adversarial"))
    def test_live_register_binding_is_certified_not_assumed(self):
        core, ov = cb.load_built(self.biz); core = json.loads(json.dumps(core))
        s = next(x for x in core["sources"]["sources"] if x["source_id"] == "S09"); s.pop("live_integration")
        rec = cb.certify(core, ov, root=self.root, d=self.biz, check_git=False); self.assertFalse(rec["checks"]["live_registers_bound"]["pass"])
        core, ov = cb.load_built(self.biz); core = json.loads(json.dumps(core)); s = next(x for x in core["sources"]["sources"] if x["source_id"] == "S09"); s["live_integration"] = "INT-NOPE"
        self.assertFalse(cb.certify(core, ov, root=self.root, d=self.biz, check_git=False)["checks"]["live_registers_bound"]["pass"])
        core, ov = cb.load_built(self.biz); core = json.loads(json.dumps(core)); s = next(x for x in core["sources"]["sources"] if x["source_id"] == "S01"); s["fingerprint_scope"] = "STRUCTURE"
        self.assertFalse(cb.certify(core, ov, root=self.root, d=self.biz, check_git=False)["checks"]["live_registers_bound"]["pass"])   # nothing but a LIVE_REGISTER may use STRUCTURE scope
    @covers(BQ, *GOV, kinds=("unit",))
    def test_real_model_declares_the_live_register(self):
        self.assertEqual(bm_sources.kind(S09), "LIVE_REGISTER"); self.assertEqual(bm_sources.scope(S09), "STRUCTURE"); self.assertEqual(S09["live_integration"], "INT-TASKS")
        self.assertEqual([s["source_id"] for s in bm_sources.SOURCES if bm_sources.scope(s) == "STRUCTURE"], ["S09"])
        snap = bb.snapshot(ROOT)["S09"]; self.assertEqual(snap["scope"], "STRUCTURE"); self.assertIsNone(snap["size"]); self.assertEqual(snap["sha256"], _struct(REAL))
        self.assertEqual(sorted(data_sync.model_source_paths()), sorted(s["path"] for s in bm_sources.SOURCES if s["currency"] == "CURRENT" and s["source_id"] != "S09"))

class L03_ProvenanceAndEvidence(unittest.TestCase):
    @covers(BQ, "daily_briefing", "deadline_management", *GOV, kinds=("unit", "completion"))
    def test_live_task_reads_carry_provenance_and_freshness(self):
        import layer, registry
        e = layer.query("INT-TASKS", "tasks.list", {}, use_cache=False)
        self.assertEqual(e["status"], "OK"); self.assertEqual(e["freshness"], "LIVE"); self.assertTrue(e["retrieved_at"]); self.assertEqual(e["authority"]["business_source"], "S09")
        self.assertEqual(registry.INTEGRATIONS["INT-TASKS"]["authority"]["business_source"], S09["source_id"]); self.assertEqual(S09["live_integration"], "INT-TASKS")
        for r in e["records"][:3]: self.assertIn("record_id", r); self.assertIn("source_updated_at", r); self.assertIn("status", r)
    @covers("audit_logging", *GOV, kinds=("unit",))
    def test_existing_verified_write_evidence_and_task_16_closure_are_intact(self):
        wc = json.loads((ROOT / ".claude" / "state" / "durable" / "write_certifications.json").read_text(encoding="utf-8"))
        c = wc["INT-TASKS"]["tasks.create"]; self.assertEqual(c["action_id"], "ACT-ef5fb6b5d0"); self.assertEqual(c["evidence"]["approved_by"], "Gev"); self.assertEqual(c["evidence"]["id"], 16)
        t = next(r for r in executors.load_tasks(REAL) if r["id"] == 16); self.assertEqual(t["status"], "Արված"); self.assertIn("ACT-ef5fb6b5d0", str(t["comment"]))
        acts = [json.loads(l) for l in (ROOT / ".claude" / "state" / "durable" / "actions.jsonl").read_text(encoding="utf-8").splitlines() if l.strip()]
        st = {a["op_id"]: a["status"] for a in acts}; self.assertEqual(st.get("ACT-ef5fb6b5d0"), "VERIFIED"); self.assertEqual(st.get("ACT-0854f55fe0"), "VERIFIED")

def _repo(name, model_sources=("Actions.md", "01_Active/Sales/Source.docx")):
    """Temp git repository shaped like the workspace (subset), with stored checksums, one commit, identity configured, no remote."""
    r = TMP / name; (r / ".claude" / "policy").mkdir(parents=True); (r / ".claude" / "state" / "durable").mkdir(parents=True); (r / ".claude" / "skills").mkdir(); (r / "00_Inbox").mkdir(); (r / "01_Active" / "Sales").mkdir(parents=True)
    shutil.copy(REAL, r / "Tasks.xlsx"); (r / "Journal.md").write_text("# J\n", encoding="utf-8"); (r / "Actions.md").write_text("# A\n", encoding="utf-8"); (r / "00_Inbox" / "Input.md").write_text("", encoding="utf-8")
    (r / "01_Active" / "Sales" / "Source.docx").write_bytes(b"src"); (r / "01_Active" / "Sales" / "Other.docx").write_bytes(b"other"); (r / ".claude" / "skills" / "x.py").write_text("x = 1\n", encoding="utf-8")
    (r / ".claude" / "state" / "durable" / "actions.jsonl").write_text("", encoding="utf-8"); (r / ".gitignore").write_text(".claude/state/*\n!.claude/state/durable/\n!.claude/state/durable/**\n", encoding="utf-8")
    shutil.copy(ROOT / ".claude" / "policy" / "workspace_policy.json", r / ".claude" / "policy" / "workspace_policy.json")
    tm.write_checksums(r, r / ".claude" / "policy" / "durable_checksums.json")
    g = lambda *a: subprocess.run(["git"] + list(a), cwd=str(r), capture_output=True, text=True, encoding="utf-8", errors="replace")
    g("init", "-q"); g("config", "user.name", "Test"); g("config", "user.email", "test@example.invalid"); g("config", "commit.gpgsign", "false"); g("add", "-A"); g("commit", "-q", "-m", "base")
    branch = g("branch", "--show-current").stdout.strip()
    return r, list(model_sources), branch

class L04_DriftClassification(unittest.TestCase):
    @covers(*GOV, DS, kinds=("unit",))
    def test_classify_path_is_deterministic(self):
        spec = tm.live_data_spec(); ms = ["Actions.md", "01_Active/Sales/Source.docx"]
        cases = {"Tasks.xlsx": "LIVE_DATA", "Journal.md": "LIVE_DATA", "00_Inbox/Input.md": "LIVE_DATA", ".claude/state/durable/actions.jsonl": "DURABLE_STATE", ".claude/policy/durable_checksums.json": "INTEGRITY_META",
                 "Actions.md": "MODEL_SOURCE", "01_Active/Sales/Source.docx": "MODEL_SOURCE", "01_Active/Sales/Other.docx": "DOCUMENT", "04_Sources/Imports/x.csv": "DOCUMENT", ".claude/skills/x.py": "PRODUCT", ".claude/policy/workspace_policy.json": "PRODUCT",
                 "CLAUDE.md": "PRODUCT", ".claude/skills/certifications/a.json": "RELEASE_ARTIFACT", ".claude/skills/registry.json": "RELEASE_ARTIFACT", "stray.txt": "UNKNOWN", "Random/thing.md": "UNKNOWN"}
        for path, want in cases.items(): self.assertEqual(tm.classify_path(path, spec, ms), want, path)
    @covers(*GOV, DS, "completion_verification", kinds=("unit", "failure", "failure_injection"))
    def test_drift_states_clean_sync_release_unclassified(self):
        r, ms, _ = _repo("drift")
        self.assertEqual(tm.classify_drift(r, ms)["state"], "CLEAN")
        _mutate_rows(r / "Tasks.xlsx"); d = tm.classify_drift(r, ms)
        self.assertEqual(d["state"], "SYNC_REQUIRED"); self.assertEqual(d["changes"], {"LIVE_DATA": ["Tasks.xlsx"]}); self.assertEqual(d["checksum_drift"], {"LIVE_DATA": ["Tasks.xlsx"]})
        (r / "Journal.md").write_text("# J\n- entry\n", encoding="utf-8"); (r / ".claude" / "state" / "durable" / "actions.jsonl").write_text("{}\n", encoding="utf-8"); (r / "01_Active" / "Sales" / "Other.docx").write_bytes(b"other2")
        d = tm.classify_drift(r, ms); self.assertEqual(d["state"], "SYNC_REQUIRED"); self.assertEqual(sorted(d["changes"]), ["DOCUMENT", "DURABLE_STATE", "LIVE_DATA"])
        (r / ".claude" / "skills" / "x.py").write_text("x = 2\n", encoding="utf-8"); d = tm.classify_drift(r, ms); self.assertEqual(d["state"], "RELEASE_REQUIRED"); self.assertEqual(d["changes"]["PRODUCT"], [".claude/skills/x.py"])
        (r / ".claude" / "skills" / "x.py").write_text("x = 1\n", encoding="utf-8"); (r / "Actions.md").write_text("# A\n- changed\n", encoding="utf-8")
        d = tm.classify_drift(r, ms); self.assertEqual(d["state"], "RELEASE_REQUIRED"); self.assertEqual(d["changes"]["MODEL_SOURCE"], ["Actions.md"])
        (r / "Actions.md").write_text("# A\n", encoding="utf-8"); (r / "stray.txt").write_text("?", encoding="utf-8")
        d = tm.classify_drift(r, ms); self.assertEqual(d["state"], "UNCLASSIFIED"); self.assertEqual(d["changes"]["UNKNOWN"], ["stray.txt"])
    @covers(*GOV, DS, kinds=("unit", "failure"))
    def test_checksum_drift_alone_is_detected_and_classified(self):
        r, ms, _ = _repo("csdrift"); cs = json.loads((r / ".claude" / "policy" / "durable_checksums.json").read_text(encoding="utf-8"))
        cs["files"]["Tasks.xlsx"]["sha256"] = "0" * 64; (r / ".claude" / "policy" / "durable_checksums.json").write_text(json.dumps(cs), encoding="utf-8")
        subprocess.run(["git", "commit", "-q", "-am", "stale checksum"], cwd=str(r), capture_output=True)
        d = tm.classify_drift(r, ms); self.assertEqual(d["state"], "SYNC_REQUIRED"); self.assertEqual(d["changes"], {}); self.assertEqual(d["checksum_drift"], {"LIVE_DATA": ["Tasks.xlsx"]})

class L05_LiveDataSync(unittest.TestCase):
    def setUp(self):
        self._env = {k: os.environ.get(k) for k in ("COMMAND_CENTER_BUSINESS_DIR", "COMMAND_CENTER_BUSINESS_ROOT")}
        os.environ["COMMAND_CENTER_BUSINESS_DIR"] = str(TMP / "no-model"); business.load(force=True)            # no built model in the temp repo → static check SKIPPED, honestly reported
    def tearDown(self):
        for k, v in self._env.items():
            if v is None: os.environ.pop(k, None)
            else: os.environ[k] = v
        business.load(force=True)
    def _head(self, r): return subprocess.run(["git", "rev-parse", "HEAD"], cwd=str(r), capture_output=True, text=True).stdout.strip()
    @covers(*GOV, DS, "task_management", "completion_verification", kinds=("unit", "completion", "failure_injection"))
    def test_sync_persists_data_only_and_leaves_the_tree_clean(self):
        r, ms, br = _repo("sync"); h0 = self._head(r)
        data_sync.model_source_paths = lambda: ms                      # the temp repo's declared model sources (restored below)
        try:
            _mutate_rows(r / "Tasks.xlsx"); (r / "Journal.md").write_text("# J\n- synced entry\n", encoding="utf-8")
            dry = data_sync.run(r, dry_run=True, push=False, log=lambda *a: None, allow_branch=br)
            self.assertEqual(dry["result"], "DRY_RUN" if False else dry["result"]); self.assertEqual(self._head(r), h0); self.assertEqual(tm.classify_drift(r, ms)["state"], "SYNC_REQUIRED")   # dry run changes nothing
            rep = data_sync.run(r, dry_run=False, push=False, log=lambda *a: None, allow_branch=br)
            self.assertEqual(rep["result"], "SYNCED_LOCAL"); self.assertNotEqual(self._head(r), h0); self.assertEqual(rep["final_state"], "CLEAN")
            steps = {s["step"]: s for s in rep["steps"]}; self.assertEqual(steps["static business model"]["status"], "SKIPPED"); self.assertEqual(steps["durable checksums"]["status"], "OK")
            cs = tm.verify_checksums(r); self.assertEqual(cs["changed"], []); self.assertEqual(cs["missing"], [])
            shown = subprocess.run(["git", "show", "--stat", "--format=%s", "HEAD"], cwd=str(r), capture_output=True, text=True, encoding="utf-8").stdout
            self.assertIn("sync: live data", shown); self.assertIn("Tasks.xlsx", shown); self.assertIn("durable_checksums.json", shown); self.assertNotIn("x.py", shown)
            self.assertEqual(tm.classify_drift(r, ms)["state"], "CLEAN")
            # recovery truth: the committed register is the mutated one
            blob = subprocess.run(["git", "show", "HEAD:Tasks.xlsx"], cwd=str(r), capture_output=True).stdout; self.assertEqual(hashlib.sha256(blob).hexdigest(), hashlib.sha256((r / "Tasks.xlsx").read_bytes()).hexdigest())
            self.assertEqual(data_sync.run(r, push=False, log=lambda *a: None, allow_branch=br)["result"], "CLEAN")
        finally: data_sync.model_source_paths = lambda: [s["path"] for s in bm_sources.SOURCES if s["currency"] == "CURRENT" and bm_sources.scope(s) == "CONTENT"]
    @covers(*GOV, DS, kinds=("unit", "failure", "adversarial"))
    def test_sync_refuses_product_model_source_unknown_and_wrong_branch(self):
        r, ms, br = _repo("refuse"); h0 = self._head(r); orig = data_sync.model_source_paths; data_sync.model_source_paths = lambda: ms
        try:
            _mutate_rows(r / "Tasks.xlsx"); (r / ".claude" / "skills" / "x.py").write_text("x = 2\n", encoding="utf-8")
            with self.assertRaises(data_sync.SyncError) as cm: data_sync.run(r, push=False, log=lambda *a: None, allow_branch=br)
            self.assertEqual(cm.exception.code, "RELEASE_REQUIRED"); self.assertEqual(self._head(r), h0)
            (r / ".claude" / "skills" / "x.py").write_text("x = 1\n", encoding="utf-8"); (r / "Actions.md").write_text("# A\n- model source edited\n", encoding="utf-8")
            with self.assertRaises(data_sync.SyncError) as cm: data_sync.run(r, push=False, log=lambda *a: None, allow_branch=br)
            self.assertEqual(cm.exception.code, "RELEASE_REQUIRED"); self.assertIn("Actions.md", cm.exception.detail); self.assertEqual(self._head(r), h0)
            (r / "Actions.md").write_text("# A\n", encoding="utf-8"); (r / "stray.txt").write_text("?", encoding="utf-8")
            with self.assertRaises(data_sync.SyncError) as cm: data_sync.run(r, push=False, log=lambda *a: None, allow_branch=br)
            self.assertEqual(cm.exception.code, "UNCLASSIFIED"); self.assertEqual(self._head(r), h0)
            (r / "stray.txt").unlink()
            with self.assertRaises(data_sync.SyncError) as cm: data_sync.run(r, push=False, log=lambda *a: None, allow_branch="another-branch")
            self.assertEqual(cm.exception.code, "WRONG_BRANCH"); self.assertEqual(self._head(r), h0)
            with self.assertRaises(data_sync.SyncError) as cm: data_sync.run(r, push=True, log=lambda *a: None, allow_branch=br)      # no upstream → nothing is claimed as pushed
            self.assertEqual(cm.exception.code, "NO_UPSTREAM")
        finally: data_sync.model_source_paths = orig
    @covers(*GOV, "audit_logging", "commitment_memory", kinds=("unit", "failure", "failure_injection"))
    def test_sync_never_exports_a_store_that_is_behind_the_versioned_durable_history(self):
        r, ms, br = _repo("regress"); h0 = self._head(r); orig = data_sync.model_source_paths; data_sync.model_source_paths = lambda: ms
        n = engine._store().count("audit") + 3                                       # the versioned history is strictly AHEAD of this process's store, whatever other suites wrote
        f = r / ".claude" / "state" / "durable" / "audit.jsonl"; f.write_text("\n".join(json.dumps({"op_id": f"hist-{i}", "recorded_at": "2026-09-01T00:00:00", "payload": {"execution_id": f"hist-{i}", "result_status": "EXECUTED"}}) for i in range(n)) + "\n", encoding="utf-8")
        tm.write_checksums(r, r / ".claude" / "policy" / "durable_checksums.json"); subprocess.run(["git", "add", "-A"], cwd=str(r)); subprocess.run(["git", "commit", "-q", "-m", "history"], cwd=str(r))
        keep = f.read_bytes(); h0 = self._head(r)
        try:
            _mutate_rows(r / "Tasks.xlsx")
            with self.assertRaises(data_sync.SyncError) as cm: data_sync.run(r, push=False, log=lambda *a: None, allow_branch=br)
            self.assertEqual(cm.exception.code, "DURABLE_REGRESSION"); self.assertIn("audit", cm.exception.detail)
            self.assertEqual(f.read_bytes(), keep, "versioned durable history must not be rewritten"); self.assertEqual(self._head(r), h0)
        finally: data_sync.model_source_paths = orig
    @covers(*GOV, kinds=("enforcement", "unit"))
    def test_gate_allows_the_sanctioned_sync_command_only_as_a_whole(self):
        gate = ROOT / ".claude" / "hooks" / "gate.py"; env = {**os.environ, "SKILL_STATE_DIR": str(TMP / "gate_state")}
        def pre(cmd):
            d = {"hook_event_name": "PreToolUse", "session_id": "live-gate", "cwd": str(ROOT), "tool_name": "Bash", "tool_input": {"command": cmd}}
            p = subprocess.run([sys.executable, str(gate), "PreToolUse"], input=json.dumps(d), capture_output=True, text=True, encoding="utf-8", env=env, cwd=str(ROOT))
            for line in p.stdout.splitlines():
                if line.startswith("{"):
                    try: return json.loads(line).get("hookSpecificOutput", {}).get("permissionDecision")
                    except ValueError: pass
            return None
        self.assertNotEqual(pre("python .claude/skills/skill.py sync --dry-run"), "deny")
        self.assertEqual(pre("python .claude/runtime/data_sync.py; rm -rf .claude/policy"), "deny")

class L06_RecoveryContract(unittest.TestCase):
    @covers(*GOV, DS, kinds=("unit",))
    def test_policy_manifest_and_gitignore_carry_the_live_data_contract(self):
        pol = tm.load_policy(); ld = pol["durability"]["live_data"]
        self.assertEqual(ld["live_data_files"], ["Tasks.xlsx", "Journal.md", "00_Inbox/Input.md"]); self.assertEqual(ld["live_state_dirs"], [".claude/state/durable"]); self.assertIn("skill.py sync", ld["sync"])
        self.assertEqual(tm.check_manifest(), []); m = tm.build(); self.assertEqual(m["policy_version"], pol["policy_version"])
        paths = {e["path"]: e for e in m["entries"]}; self.assertEqual(paths["Tasks.xlsx"]["durability"], "VERSION_DIRECTLY"); self.assertEqual(paths[".claude/state/durable/actions.jsonl"]["durability"], "VERSION_DIRECTLY")
        self.assertIn("Tasks.xlsx", tm.durable_files(ROOT)); self.assertIn("Tasks.xlsx", json.loads((ROOT / ".claude" / "policy" / "durable_checksums.json").read_text(encoding="utf-8"))["files"])
    @covers(*GOV, kinds=("unit",))
    def test_real_workspace_drift_is_reported_never_hidden(self):
        d = data_sync.plan(ROOT); self.assertIn(d["state"], tm.DRIFT_STATES); self.assertIsInstance(d["changes"], dict)
        if d["state"] == "CLEAN": self.assertEqual(d["changes"], {}); self.assertEqual(d["checksum_drift"], {}); self.assertFalse(d["ahead"])

if __name__ == "__main__":
    unittest.main(verbosity=2)
