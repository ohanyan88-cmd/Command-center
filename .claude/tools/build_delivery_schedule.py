# -*- coding: utf-8 -*-
"""Կառուցում է Գադուկյանին ուղարկվող գրաֆիկը՝ ուղիղ ռեեստրից։
Դուրս է գնում միայն ext=True կետերը՝ deliv դաշտի ձևակերպմամբ։
Ելք՝ 01_Active/Operations/Delivery-schedule-<date>.xlsx  և  Delivery-schedule-whatsapp-<date>.txt"""
import sys, os, io, datetime
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border

from xlsx_style import *
from task_registry_data import IN, OUT, ANSWERS, TODAY

FAR = datetime.date(2099, 1, 1)
def is_open(x): return x["st"] in OPEN_ST
def shows(x):   return x["ext"] and is_open(x) and x["deliv"]

mine   = sorted([x for x in IN if shows(x) and x["way"] == "ես" and x["due"]],
                key=lambda x: (x["due"], x["id"]))
theirs = sorted([x for x in IN if shows(x) and x["way"] == "նա"],
                key=lambda x: (x["due"] or FAR, x["id"]))
team   = sorted([x for x in OUT if shows(x)],
                key=lambda x: ((x["due"] or x["nxt"] or FAR), x["to"]))

def when(d):
    return f"{d.strftime('%m-%d')} · {HY[d.weekday()]}" if d else "ժամկետ չկա"

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
DEST = os.path.join(ROOT, "01_Active", "Operations")
stamp = TODAY.strftime("%Y-%m-%d")
RULE = "─────────────────────"

# ══════════════════════ WhatsApp-ի տեքստ ══════════════════════
ready = [x for x in ANSWERS if x["ready"] and x["a"]]

L = ["Գև ջան, ահա պատասխանները և գրաֆիկը։", ""]
if ready:
    L += [RULE, "", "*ՊԱՏԱՍԽԱՆՆԵՐ*", ""]
    for x in ready:
        L.append(f"*{x['q']}*")
        for line in x["a"]:
            L.append(line)
        L.append("")
L += [RULE, "", "*ԳՐԱՖԻԿ — ԵՍ ԵՄ ՏԱԼՈՒ*", "",
      "_Ամսաթվերի մի մասը քո նշածն է, մի մասը՝ իմ առաջարկը — "
      "գրված է ամեն կետի տակ։_", ""]
prev = None
for x in mine:
    if x["due"] != prev:
        L.append(f"*{when(x['due'])}*")
        prev = x["due"]
    L += [f"• {x['deliv']}", f"  _{x['src']}_", ""]

if theirs:
    L += [RULE, "", "*ՔԵԶՆԻՑ ԵՄ ՍՊԱՍՈՒՄ*", ""]
    for x in theirs:
        L += [f"• {x['deliv']}", f"  _{when(x['due'])}_", ""]

if team:
    L += [RULE, "", "*ԹԻՄԻՑ ԵՄ ՍՊԱՍՈՒՄ*", ""]
    for x in team:
        L += [f"• *{x['to']}* — {x['deliv']}",
              f"  _{when(x['due'] or x['nxt'])}_", ""]

L += [RULE, "",
      "Եթե որևէ ամսաթիվ չի սազում կամ մի բան բաց եմ թողել՝ ասա, կուղղեմ։"]
io.open(os.path.join(DEST, f"Delivery-schedule-whatsapp-{stamp}.txt"), "w",
        encoding="utf-8").write("\n".join(L))

# ══════════════════════ Excel ══════════════════════
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Գրաֆիկ"
ws.sheet_view.showGridLines = False
for k, v in {"A":2,"B":2.4,"C":17,"D":68,"E":22,"F":24,"G":2}.items():
    ws.column_dimensions[k].width = v

def paint2(r1,c1,r2,c2,color):
    f = PatternFill("solid", fgColor=color)
    for r in range(r1,r2+1):
        for c in range(c1,c2+1): ws.cell(row=r,column=c).fill = f

def out2(r1,c1,r2,c2,color=LINE):
    for r in range(r1,r2+1):
        for c in range(c1,c2+1):
            ws.cell(row=r,column=c).border = Border(
                top=side(color) if r==r1 else None,
                bottom=side(color) if r==r2 else None,
                left=side(color) if c==c1 else None,
                right=side(color) if c==c2 else None)

def put2(r,c1,c2,t,*,size=10,bold=False,italic=False,color=INK,
         align="left",wrap=False,h=None):
    if c2>c1: ws.merge_cells(start_row=r,start_column=c1,end_row=r,end_column=c2)
    cell = ws.cell(row=r,column=c1,value=t)
    cell.font = Font(name="Calibri",size=size,bold=bold,italic=italic,color=color)
    cell.alignment = Alignment(horizontal=align,vertical="center",wrap_text=wrap)
    if h: ws.row_dimensions[r].height = h
    return cell

paint2(1, 1, (len(mine)+len(theirs)+len(team))*2 + 44, 7, PAPER)
put2(1,3,6,"ԳՐԱՖԻԿ — ԻՆՉ ԵՐԲ Է ՊԱՏՐԱՍՏ",size=17,bold=True,color=NAVY,h=36)
put2(2,3,6,f"{stamp}, {HY[TODAY.weekday()]}   ·   "
     "«Ժամկետը» սյունակը ցույց է տալիս՝ ամսաթիվը դու ես նշել, թե ես եմ առաջարկում",
     size=9,color=MUTED,h=17)
paint2(3,2,3,6,NAVY); ws.row_dimensions[3].height = 2.5
ws.row_dimensions[4].height = 12

def block(r, title, color, third):
    put2(r,3,6,title,size=12,bold=True,color=color,h=26)
    paint2(r+1,3,r+1,6,color); ws.row_dimensions[r+1].height = 1.5
    ws.row_dimensions[r+2].height = 6
    r += 3
    for j,h in enumerate(["Ամսաթիվ","Ինչ",third,"Ժամկետը"],3):
        cell = ws.cell(row=r,column=j,value=h)
        cell.fill = PatternFill("solid",fgColor=NAVY)
        cell.font = Font(name="Calibri",color="FFFFFF",bold=True,size=9)
        cell.alignment = Alignment(horizontal="center",vertical="center")
    ws.row_dimensions[r].height = 22
    return r + 1

def row(r, a, b, c, d, accent=INK, bg=CARD):
    paint2(r,3,r,6,bg); out2(r,3,r,6)
    put2(r,3,3,a,size=9.5,bold=True,color=accent)
    put2(r,4,4,b,size=9.5,color=INK,wrap=True)
    put2(r,5,5,c,size=9,color=MUTED,wrap=True)
    put2(r,6,6,d,size=9,color=(RED if d.startswith("Իմ") else MUTED),
         wrap=True,h=30)
    return r + 1

r = block(5, "ԵՍ ԵՄ ՏԱԼՈՒ", NAVY, "")
prev = None
for x in mine:
    n = (x["due"] - TODAY).days
    accent = RED if n < 0 else AMBER if n == 0 else INK
    bg = RED_BG if n < 0 else AMBER_BG if n == 0 else CARD
    r = row(r, when(x["due"]) if x["due"] != prev else "", x["deliv"], "",
            x["src"], accent, bg)
    prev = x["due"]

if theirs:
    r = block(r + 1, "ՔԵԶՆԻՑ ԵՄ ՍՊԱՍՈՒՄ", AMBER, "")
    for x in theirs:
        r = row(r, when(x["due"]), x["deliv"], "", x["src"])

if team:
    r = block(r + 1, "ԹԻՄԻՑ ԵՄ ՍՊԱՍՈՒՄ", SLATE, "Ումից")
    for x in team:
        r = row(r, when(x["due"] or x["nxt"]), x["deliv"], x["to"], x["src"])

r += 2
put2(r,3,6,"Եթե որևէ ամսաթիվ չի սազում կամ մի բան բաց եմ թողել՝ ասա, կուղղեմ։",
     size=9,italic=True,color=MUTED,h=18)
wb.save(os.path.join(DEST, f"Delivery-schedule-{stamp}.xlsx"))

print(f"Գրաֆիկ_{stamp}  ·  ես եմ տալու՝ {len(mine)}  ·  "
      f"քեզնից՝ {len(theirs)}  ·  թիմից՝ {len(team)}  →  01_Active/Operations")
