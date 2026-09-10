# -*- coding: utf-8 -*-
"""Ընդհանուր ոճ և ցանց՝ բոլոր թերթերի համար։
Մեկ ցանց՝ A լուսանցք · B շեշտ · C..J բովանդակություն (8 սյունակ) · K լուսանցք։
8 սյունակը բաժանվում է 2-ի, 4-ի կամ 8-ի՝ միշտ սիմետրիկ։"""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ── գունապնակ ──
NAVY   = "1F3864"
INK    = "24292F"
MUTED  = "8B9099"
FAINT  = "B9BEC6"
LINE   = "E1E5EA"
PAPER  = "F5F7F9"
CARD   = "FFFFFF"
RED    = "B42318"; RED_BG    = "FDECEA"
AMBER  = "9A6A00"; AMBER_BG  = "FDF5E4"
BLUE   = "18539E"; BLUE_BG   = "EBF2FA"
GREEN  = "2F6B36"; GREEN_BG  = "EDF6EE"
SLATE  = "8B9099"; SLATE_BG  = "F0F2F4"

PRI = {"Բարձր": RED, "Միջին": AMBER, "Ցածր": FAINT}
ST  = {"Նոր":       (BLUE,  BLUE_BG),
       "Ընթացքում": (AMBER, AMBER_BG),
       "Սպասում":   (SLATE, SLATE_BG),
       "Կատարված":  (GREEN, GREEN_BG),
       "Չեղարկված": (FAINT, SLATE_BG)}
OPEN_ST   = ("Նոր", "Ընթացքում", "Սպասում")
CLOSED_ST = ("Կատարված", "Չեղարկված")

STATUSES   = '"Նոր,Ընթացքում,Սպասում,Կատարված,Չեղարկված"'
PRIORITIES = '"Բարձր,Միջին,Ցածր"'

HY = ["երկուշաբթի","երեքշաբթի","չորեքշաբթի","հինգշաբթի",
      "ուրբաթ","շաբաթ","կիրակի"]
HY_SHORT = ["երկ","երք","չրք","հնգ","ուրբ","շբթ","կիր"]

# ── ցանց ──
MARGIN_L, STRIPE, C_FIRST, C_LAST, MARGIN_R = 1, 2, 3, 10, 11
COL_W = 15.5

def grid(ws, stripe=True):
    """Դնում է միասնական սյունակների լայնությունները։"""
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 2.4 if stripe else 2
    for i in range(C_FIRST, C_LAST + 1):
        ws.column_dimensions[chr(64 + i)].width = COL_W
    ws.column_dimensions["K"].width = 2

def half(n):
    """8 սյունակը՝ n հավասար մասի։ Վերադարձնում է (սկիզբ, վերջ) զույգեր։"""
    step = 8 // n
    return [(C_FIRST + i * step, C_FIRST + (i + 1) * step - 1) for i in range(n)]

def side(c=LINE, style="thin"): return Side(style=style, color=c)

def paint(ws, r1, c1, r2, c2, color):
    f = PatternFill("solid", fgColor=color)
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            ws.cell(row=r, column=c).fill = f

def outline(ws, r1, c1, r2, c2, color=LINE, style="thin"):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            ws.cell(row=r, column=c).border = Border(
                top=side(color, style)    if r == r1 else None,
                bottom=side(color, style) if r == r2 else None,
                left=side(color, style)   if c == c1 else None,
                right=side(color, style)  if c == c2 else None)

def put(ws, r, c1, c2, text, *, size=10, bold=False, italic=False,
        color=INK, align="left", wrap=False, h=None):
    if c2 > c1:
        ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
    cell = ws.cell(row=r, column=c1, value=text)
    cell.font = Font(name="Calibri", size=size, bold=bold,
                     italic=italic, color=color)
    cell.alignment = Alignment(horizontal=align, vertical="center",
                               wrap_text=wrap)
    if h: ws.row_dimensions[r].height = h
    return cell

def title_block(ws, title, subtitle):
    put(ws, 1, C_FIRST, C_LAST, title, size=17, bold=True, color=NAVY, h=36)
    put(ws, 2, C_FIRST, C_LAST, subtitle, size=9, color=MUTED, h=17)
    paint(ws, 3, STRIPE, 3, C_LAST, NAVY)
    ws.row_dimensions[3].height = 2.5
    ws.row_dimensions[4].height = 12
    return 5

def section(ws, r, text, color, right=""):
    put(ws, r, C_FIRST, C_LAST - 2, text, size=12, bold=True, color=color, h=26)
    if right:
        put(ws, r, C_LAST - 1, C_LAST, right, size=9, bold=True,
            color=color, align="right")
    paint(ws, r + 1, C_FIRST, r + 1, C_LAST, color)
    ws.row_dimensions[r + 1].height = 1.5
    ws.row_dimensions[r + 2].height = 6
    return r + 3

def date_badge(due, st, today, with_time=""):
    """Վերադարձնում է (տեքստ, գույն)։"""
    if st in CLOSED_ST:  return ("ՓԱԿՎԱԾ", GREEN)
    if due is None:      return ("ժամկետ չկա", FAINT)
    n = (due - today).days
    d = due.strftime("%m-%d") + (f" {with_time}" if with_time else "")
    if n < 0:  return (f"ԺԱՄԿԵՏԱՆՑ · {d} · {-n} օր", RED)
    if n == 0: return (f"ԱՅՍՕՐ · {d}", AMBER)
    if n == 1: return (f"ՎԱՂԸ · {d} · {HY[due.weekday()]}", AMBER)
    return (f"{d} · {HY_SHORT[due.weekday()]} · {n} օր", MUTED)

def header_row(ws, r, cols, first_col=1):
    for j, (h, w) in enumerate(cols, first_col):
        c = ws.cell(row=r, column=j, value=h)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.font = Font(name="Calibri", color="FFFFFF", bold=True, size=9.5)
        c.alignment = Alignment(horizontal="center", vertical="center",
                                wrap_text=True)
        ws.column_dimensions[chr(64 + j)].width = w
    ws.row_dimensions[r].height = 34
